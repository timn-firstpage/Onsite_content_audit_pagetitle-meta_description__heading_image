import importlib.util
import csv
import json
import os
import sys
import subprocess
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'scripts'))


class ContentTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.path = Path(self.tmp.name)

    def write(self, name, headers, rows):
        with (self.path / name).open('w', encoding='utf-8-sig', newline='') as f:
            w = csv.writer(f)
            w.writerow(headers)
            w.writerows(rows)

    def full_exports(self):
        self.write('page_titles_all.csv', ['Address', 'Indexability', 'Title 1', 'Title 1 Pixel Width'],
                   [['https://example.com/a', 'Indexable', 'Title', '580']])
        self.write('meta_description_all.csv', ['Address', 'Indexability', 'Meta Description 1', 'Meta Description 1 Pixel Width'],
                   [['https://example.com/a', 'Indexable', 'description', '880']])
        for name, headers in {
            'page_titles_missing.csv': ['Address'], 'page_titles_duplicate.csv': ['Address', 'Title 1'],
            'page_titles_multiple.csv': ['Address', 'Title 1', 'Title 2'],
            'meta_description_missing.csv': ['Address'], 'meta_description_duplicate.csv': ['Address', 'Meta Description 1'],
            'h1_missing.csv': ['Address'], 'h1_duplicate.csv': ['Address', 'H1-1', 'H1-2'],
            'h1_multiple.csv': ['Address', 'H1-1', 'H1-2'],
            'images_missing_alt_attribute_and_text_inlinks.csv': ['Source', 'Destination'],
        }.items():
            self.write(name, headers, [])
        self.write('images_all.csv', ['Address'], [['https://cdn.example.com/a.png']])
        self.write('all_image_inlinks.csv', ['Source', 'Destination', 'Size (Bytes)'],
                   [['https://example.com/a', 'https://cdn.example.com/a.png', '204800']])

    def audit(self, **kwargs):
        if not (ROOT / 'scripts' / 'extract.py').exists():
            self.fail('Content extraction interface has not been implemented')
        import extract
        return extract.audit(self.path, **kwargs)

    def check(self, data, item):
        return next(c for c in data['checks'] if c['id'] == item)

    def test_native_empty_exports_pass_and_scope_is_twelve(self):
        self.full_exports()
        data = self.audit()
        self.assertEqual(len(data['checks']), 12)
        self.assertEqual({c['result'] for c in data['checks']}, {'√'})

    def test_native_duplicates_not_recomputed_from_all(self):
        self.full_exports()
        self.write('page_titles_all.csv', ['Address', 'Indexability', 'Title 1', 'Title 1 Pixel Width'],
                   [['https://example.com/a', 'Indexable', 'Same', 100], ['https://example.com/b', 'Non-Indexable', 'Same', 100]])
        self.assertEqual(self.check(self.audit(), '17.2')['result'], '√')

    def test_h1_drops_only_pure_same_page_collision(self):
        self.full_exports()
        self.write('h1_duplicate.csv', ['Address', 'H1-1', 'H1-2'],
                   [['a', 'Unique', 'Unique'], ['b', 'Shared', 'Shared'], ['c', 'Shared', 'Other']])
        self.assertEqual(self.check(self.audit(), '19.2')['row_count'], 2)

    def test_thresholds_indexability_and_description_characters(self):
        self.full_exports()
        self.write('page_titles_all.csv', ['Address', 'Indexability', 'Title 1', 'Title 1 Pixel Width'],
                   [['a', 'Indexable', 'Long', 581], ['b', 'Non-Indexable', 'Long', 900]])
        self.write('meta_description_all.csv', ['Address', 'Indexability', 'Meta Description 1', 'Meta Description 1 Pixel Width'],
                   [['a', 'Indexable', 'x' * 156, 100], ['b', 'Indexable', 'short', 881]])
        self.assertEqual(self.check(self.audit(), '17.3')['row_count'], 1)
        self.assertEqual(self.check(self.audit(), '18.3')['rows'][0][0], 'a')

    def test_exact_targets_only_exclude_missing_title_description(self):
        self.full_exports()
        self.write('page_titles_missing.csv', ['Address'], [['https://example.com/a'], ['https://example.com/a?x=1']])
        self.write('meta_description_missing.csv', ['Address'], [['https://example.com/a']])
        self.write('h1_missing.csv', ['Address'], [['https://example.com/a']])
        data = self.audit(targets=['https://example.com/a'])
        self.assertEqual(self.check(data, '17.1')['row_count'], 1)
        self.assertEqual(self.check(data, '18.1')['row_count'], 0)
        self.assertEqual(self.check(data, '19.1')['row_count'], 1)

    def test_missing_and_invalid_exports_are_unknown_not_clean(self):
        self.full_exports()
        (self.path / 'h1_missing.csv').unlink()
        self.write('page_titles_all.csv', ['Address', 'Indexability', 'Title 1'], [['a', 'Indexable', 'Title']])
        data = self.audit()
        self.assertEqual(self.check(data, '19.1')['result'], 'Human Check')
        self.assertEqual(self.check(data, '17.3')['result'], 'Human Check')

    def test_known_oversize_survives_partial_cdn_size_evidence(self):
        self.full_exports()
        self.write('images_all.csv', ['Address'], [])
        self.write('all_image_inlinks.csv', ['Source', 'Destination', 'Size (Bytes)'],
                   [['a', 'https://cdn.example.com/big.png', 204801], ['b', 'https://cdn.example.com/big.png', 204801],
                    ['a', 'https://cdn.example.com/unknown.png', '']])
        c = self.check(self.audit(), '20.2')
        self.assertEqual((c['result'], c['row_count'], c['unique_assets']), ('X', 2, 1))
        self.assertTrue(c['gaps'])

    def test_images_without_capture_or_missing_alt_report_unknown(self):
        self.full_exports()
        (self.path / 'images_missing_alt_attribute_and_text_inlinks.csv').unlink()
        self.assertEqual(self.check(self.audit(), '20.1')['result'], 'Human Check')
        self.write('images_all.csv', ['Address'], [])
        self.write('all_image_inlinks.csv', ['Source', 'Destination', 'Size (Bytes)'], [])
        self.assertEqual(self.check(self.audit(), '20.2')['result'], 'Human Check')

    def test_invalid_numbers_do_not_prove_pass(self):
        self.full_exports()
        self.write('all_image_inlinks.csv', ['Source', 'Destination', 'Size (Bytes)'], [['a', 'img', 'bad']])
        self.assertEqual(self.check(self.audit(), '20.2')['result'], 'Human Check')

    def test_captured_asset_missing_from_inlinks_is_unknown(self):
        self.full_exports()
        self.write('images_all.csv', ['Address'], [['https://cdn.example.com/a.png'], ['https://cdn.example.com/b.png']])
        self.assertEqual(self.check(self.audit(), '20.2')['result'], 'Human Check')

    def test_header_only_description_all_cannot_prove_clean(self):
        self.full_exports()
        self.write('meta_description_all.csv', ['Address', 'Indexability', 'Meta Description 1'], [])
        self.assertEqual(self.check(self.audit(), '18.3')['result'], 'Human Check')

    def test_empty_crawl_is_rejected(self):
        self.full_exports()
        self.write('page_titles_all.csv', ['Address', 'Indexability', 'Title 1 Pixel Width'], [])
        with self.assertRaisesRegex(ValueError, 'Page Titles:All'):
            self.audit()

    def test_ambiguous_report_is_not_silently_selected(self):
        self.full_exports()
        self.write('old_h1_missing.csv', ['Address'], [])
        self.assertEqual(self.check(self.audit(), '19.1')['result'], 'Human Check')

    def test_report_exact_tabs_literal_values_and_no_overwrite(self):
        self.full_exports()
        self.write('page_titles_duplicate.csv', ['Address', 'Title 1'], [['a', '=1+1'], ['b', '=1+1']])
        data = self.audit()
        import report
        data.update(site_name='Example', site_url='https://example.com/', audit_date='2026-10-12')
        outputs = report.build_report(data, self.path / 'reports')
        self.assertEqual([p.name for p in outputs], [
            '17.1-17.4 Example_page_titles_audit_2026-10-12.xlsx', '18.1-18.3 Example_meta_descriptions_audit_2026-10-12.xlsx',
            '19.1-19.3 Example_headings_audit_2026-10-12.xlsx', '20.1-20.2 Example_images_audit_2026-10-12.xlsx'])
        import openpyxl
        wb = openpyxl.load_workbook(outputs[0])
        self.addCleanup(wb.close)
        self.assertEqual(wb.sheetnames, ['Checklist', '17.2 Page title - Duplicate'])
        self.assertEqual(wb.sheetnames[0], 'Checklist')
        headers = [c.value for c in wb['17.2 Page title - Duplicate'][1]]
        self.assertEqual(headers, ['Address', 'Titles', 'New Title', 'Instructions (If have)'])
        value = wb['17.2 Page title - Duplicate'].cell(2, headers.index('Titles') + 1)
        self.assertEqual(value.value, '=1+1')
        self.assertEqual(value.data_type, 's')
        self.assertNotIn('20.2 Image - Oversize', wb.sheetnames)
        with self.assertRaises(FileExistsError):
            report.build_report(data, self.path / 'reports')

    def test_report_clean_and_unknown_checks_stay_on_checklist(self):
        self.full_exports()
        import report
        import openpyxl
        for unknown in (False, True):
            if unknown:
                (self.path / 'h1_missing.csv').unlink()
            data = self.audit()
            data.update(site_name='Example', site_url='https://example.com/', audit_date='2026-10-12')
            outputs = report.build_report(data, self.path / str(unknown))
            self.assertEqual(len(outputs), 4)
            for output, count in zip(outputs, (4, 3, 3, 2)):
                wb = openpyxl.load_workbook(output)
                try:
                    self.assertEqual(wb.sheetnames, ['Checklist'])
                    self.assertEqual(wb['Checklist'].max_row, count + 1)
                    self.assertEqual(len(data['checks']), 12)
                    self.assertTrue(all(not c['rows'] for c in data['checks']))
                    if output == outputs[2]:
                        self.assertEqual(wb['Checklist']['C2'].value, 'Human Check' if unknown else '√')
                finally:
                    wb.close()

    def test_report_merges_x_rows_by_header_without_changing_json(self):
        self.full_exports()
        self.write('page_titles_missing.csv', ['Address'], [['https://example.com/missing']])
        self.write('page_titles_all.csv', ['Address', 'Indexability', 'Title 1', 'Title 1 Pixel Width'],
                   [['https://example.com/long', 'Indexable', 'Long', 581]])
        self.write('images_missing_alt_attribute_and_text_inlinks.csv', ['Source', 'Destination'],
                   [['https://example.com/a', 'https://cdn.example.com/a.png']])
        self.write('all_image_inlinks.csv', ['Source', 'Destination', 'Size (Bytes)'],
                   [['https://example.com/a', 'https://cdn.example.com/a.png', 204801]])
        data = self.audit()
        data.update(site_name='Example', site_url='https://example.com/', audit_date='2026-10-13')
        before = json.dumps(data, sort_keys=True)
        import report
        import openpyxl
        outputs = report.build_report(data, self.path / 'merged')
        self.assertEqual(json.dumps(data, sort_keys=True), before)
        wb = openpyxl.load_workbook(outputs[0])
        try:
            self.assertEqual(wb.sheetnames, ['Checklist', '17.1 Page title -Missing', '17.3 Page title -Too Long'])
            self.assertEqual(wb['17.1 Page title -Missing']['A2'].value, 'https://example.com/missing')
            self.assertEqual(wb['17.3 Page title -Too Long']['C2'].value, '581')
        finally:
            wb.close()
        wb = openpyxl.load_workbook(outputs[3])
        try:
            self.assertEqual(wb.sheetnames, ['Checklist', '20.1 Image - Alt Text Missing', '20.2 Image - Oversize'])
            self.assertEqual(wb['20.1 Image - Alt Text Missing']['A2'].value, 'https://example.com/a')
            self.assertEqual(wb['20.2 Image - Oversize']['C2'].value, '204801')
        finally:
            wb.close()

    def test_template_detail_names_and_columns_for_18_to_20(self):
        self.full_exports()
        data = self.audit()
        expected = {
            '18.1': ('18. Descriptions - Missing', ['Address', 'Remarks'], ['missing', 'note']),
            '18.2': ('18.2 Descriptions - Duplicate', ['Address', 'Meta Description'], ['dup', 'Description']),
            '18.3': ('18.3 Descriptions - Too Long', ['Address', 'Meta Description', 'Pixel Width'], ['long', 'Long description', 156]),
            '19.1': ('19.1 H1 - Missing', ['Address', 'Remarks'], ['missing', 'note']),
            '19.2': ('19.2 H1 - Duplicate', ['Address', 'H1', 'Instructions'], ['dup', 'Heading', 'note']),
            '19.3': ('19.3 H1 - Multiple', ['Address', 'H1-1', 'H1-2'], ['multiple', 'First', 'Second']),
            '20.1': ('20.1 Image - Alt Text Missing', ['Source', 'Destination', 'Example Screenshots (If Applicable)'],
                     ['page', 'image.png', 'screenshot reference']),
            '20.2': ('20.2 Image - Oversize', ['Source', 'Destination', 'Size (Bytes)'], ['page', 'image.png', 204801]),
        }
        for item, (_, _, row) in expected.items():
            self.check(data, item).update(result='X', rows=[row], row_count=1)
        self.check(data, '18.3')['pixel_widths'] = [901]
        data.update(site_name='Example', site_url='https://example.com/', audit_date='2026-10-15')
        original = json.dumps(data, sort_keys=True)
        import report
        import openpyxl
        outputs = report.build_report(data, self.path / 'template')
        for index, prefix in ((1, '18.'), (2, '19.'), (3, '20.')):
            items = [item for item in expected if item.startswith(prefix)]
            wb = openpyxl.load_workbook(outputs[index])
            try:
                self.assertEqual(wb.sheetnames, ['Checklist'] + [expected[item][0] for item in items])
                for item in items:
                    sheet, headers, row = expected[item]
                    self.assertEqual([c.value for c in wb[sheet][1]], headers)
                    wanted = list(row)
                    if item == '18.3':
                        wanted[2] = 901
                    self.assertEqual([c.value for c in wb[sheet][2]], [str(v) for v in wanted])
            finally:
                wb.close()
        self.assertEqual(json.dumps(data, sort_keys=True), original)

    def test_description_pixel_display_is_not_character_judgment(self):
        self.full_exports()
        self.write('meta_description_all.csv', ['Address', 'Indexability', 'Meta Description 1', 'Meta Description 1 Pixel Width'],
                   [['short', 'Indexable', 'x' * 100, 1200], ['longer', 'Indexable', 'x' * 160, 200],
                    ['long', 'Indexable', 'x' * 156, 901]])
        data = self.audit()
        check = self.check(data, '18.3')
        self.assertEqual([r[0] for r in check['rows']], ['longer', 'long'])
        self.assertEqual([r[2] for r in check['rows']], [160, 156])
        self.assertEqual(check['pixel_widths'], [200, 901])

    def test_missing_pixel_display_stays_blank_not_character_count(self):
        self.full_exports()
        self.write('meta_description_all.csv', ['Address', 'Indexability', 'Meta Description 1'],
                   [['long', 'Indexable', 'x' * 156]])
        data = self.audit()
        check = self.check(data, '18.3')
        self.assertEqual(check['pixel_widths'], [None])
        self.assertEqual(check['result'], 'X')
        self.assertFalse(check['gaps'])
        self.assertTrue(check['display_gaps'])
        data.update(site_name='Example', site_url='https://example.com/', audit_date='2026-10-15')
        import report
        import openpyxl
        paths = report.build_report(data, self.path / 'missing-pixels')
        wb = openpyxl.load_workbook(paths[1])
        try:
            self.assertEqual(wb['18.3 Descriptions - Too Long']['C1'].value, 'Pixel Width')
            self.assertIsNone(wb['18.3 Descriptions - Too Long']['C2'].value)
        finally:
            wb.close()

    def test_title_four_detail_schemas_preserve_json(self):
        self.full_exports()
        self.write('page_titles_missing.csv', ['Address'], [['missing']])
        self.write('page_titles_duplicate.csv', ['Address', 'Title 1'], [['duplicate', 'Same']])
        self.write('page_titles_multiple.csv', ['Address', 'Title 1', 'Title 2'], [['multiple', 'First', 'Second']])
        self.write('page_titles_all.csv', ['Address', 'Indexability', 'Title 1', 'Title 1 Pixel Width'],
                   [['long', 'Indexable', 'Long', 581]])
        data = self.audit()
        data.update(site_name='Example', site_url='https://example.com/', audit_date='2026-10-13')
        self.check(data, '17.1')['rows'][0] += ['Replacement', 'Missing remarks']
        self.check(data, '17.2')['rows'][0] += ['Unique replacement', 'Duplicate instructions']
        self.check(data, '17.3')['rows'][0] += ['Short replacement']
        self.check(data, '17.4')['rows'][0] += ['Keep one title']
        original = json.dumps(data, sort_keys=True)
        import report
        import openpyxl
        outputs = report.build_report(data, self.path / 'three-tabs')
        wb = openpyxl.load_workbook(outputs[0])
        try:
            expected = {
                '17.1 Page title -Missing': (['Address', 'New Title', 'Remarks'], ['missing', 'Replacement', 'Missing remarks']),
                '17.2 Page title - Duplicate': (['Address', 'Titles', 'New Title', 'Instructions (If have)'],
                                               ['duplicate', 'Same', 'Unique replacement', 'Duplicate instructions']),
                '17.3 Page title -Too Long': (['Address', 'Title - 1', 'Pixel Width ( max 580 pixels)', 'Revised Title'],
                                             ['long', 'Long', '581', 'Short replacement']),
                '17.4 Page title - Multiple': (['Address', 'Title - 1', 'Title - 2', 'Instructions'],
                                              ['multiple', 'First', 'Second', 'Keep one title']),
            }
            self.assertEqual(wb.sheetnames, ['Checklist'] + list(expected))
            for sheet, rows in expected.items():
                self.assertEqual([[c.value for c in row] for row in wb[sheet]], list(rows))
            self.assertEqual(wb['Checklist']['C5'].value, 'X')
        finally:
            wb.close()
        self.assertEqual(json.dumps(data, sort_keys=True), original)
        self.assertEqual(self.check(data, '17.4')['rows'], [['multiple', 'First', 'Second', 'Keep one title']])

    def test_report_existing_any_group_prevents_partial_output(self):
        self.full_exports()
        data = self.audit()
        data.update(site_name='Example', site_url='https://example.com/', audit_date='2026-10-13')
        destination = self.path / 'conflict'
        destination.mkdir()
        existing = destination / '20.1-20.2 Example_images_audit_2026-10-13.xlsx'
        existing.write_bytes(b'original')
        import report
        with self.assertRaises(FileExistsError):
            report.build_report(data, destination)
        self.assertEqual(list(destination.iterdir()), [existing])
        self.assertEqual(existing.read_bytes(), b'original')

    def test_report_verification_failure_removes_only_new_reports(self):
        self.full_exports()
        data = self.audit()
        data.update(site_name='Example', site_url='https://example.com/', audit_date='2026-10-13')
        destination = self.path / 'failure'
        destination.mkdir()
        unrelated = destination / 'keep.txt'
        unrelated.write_text('keep')
        import report
        verify = report.verify_workbook
        def fail_second_file(source, expected):
            if isinstance(source, Path) and 'meta_descriptions' in source.name:
                raise ValueError('simulated verification failure')
            return verify(source, expected)
        with patch.object(report, 'verify_workbook', side_effect=fail_second_file):
            with self.assertRaisesRegex(ValueError, 'simulated'):
                report.build_report(data, destination)
        self.assertEqual(list(destination.iterdir()), [unrelated])

    def cli(self, config, run='run'):
        script = ROOT / 'scripts' / 'onsite_checklist.py'
        self.assertTrue(script.exists(), 'Config/CLI interface has not been implemented')
        return subprocess.run([sys.executable, str(script), '--config', str(config), '--run-dir', str(self.path / run)],
                              capture_output=True, text=True, encoding='utf-8')

    def config(self, site='https://example.com/', extra=None):
        data = {'site': {'name': 'Example', 'start_url': site},
                'source': {'mode': 'exports', 'export_directory': '.'},
                'output': {'audit_date': '2026-10-12', 'timezone': 'Asia/Hong_Kong', 'language': 'en'}}
        data.update(extra or {})
        path = self.path / 'config.json'
        path.write_text(json.dumps(data), encoding='utf-8')
        return path

    def test_cli_shared_config_relative_paths_and_no_overwrite(self):
        self.full_exports()
        config = self.config()
        result = self.cli(config)
        self.assertEqual(result.returncode, 0, result.stderr)
        data = json.loads((self.path / 'run' / 'content' / 'findings.json').read_text(encoding='utf-8'))
        self.assertEqual(data['audit_date'], '2026-10-12')
        self.assertTrue((self.path / 'run' / 'content' / 'manifest.json').exists())
        status = json.loads((self.path / 'run' / 'content' / 'run-status.json').read_text())
        self.assertEqual(status['state'], 'complete')
        self.assertEqual(len(status['reports']), 4)
        self.assertTrue(all(Path(p).exists() for p in status['reports']))
        self.assertEqual(len(list((self.path / 'run' / 'content').glob('*.xlsx'))), 4)
        self.assertNotIn('report', status)
        self.assertNotEqual(self.cli(config).returncode, 0)

    def test_cli_global_config_handover_and_shared_usage_preserved(self):
        self.full_exports()
        config = self.config(extra={'checks': {'robots': True}, 'budget': {'max_mcp_calls': 7}})
        original = config.read_bytes()
        run = self.path / 'shared-run'
        run.mkdir()
        usage = run / 'usage.json'
        usage.write_text('{"mcp_calls": 3}', encoding='utf-8')
        handover = run / 'sf-handover.json'
        handover.write_text(json.dumps({'site_url': 'https://example.com/', 'crawl': {'id': 'context-only'},
                                       'gaps': ['Partial crawl']}), encoding='utf-8')
        result = subprocess.run([sys.executable, str(ROOT / 'scripts' / 'onsite_checklist.py'), '--run-dir', str(run)],
                                env={**os.environ, 'ONSITE_AUDIT_CONFIG': str(config)},
                                capture_output=True, text=True, encoding='utf-8')
        self.assertEqual(result.returncode, 0, result.stderr)
        manifest = json.loads((run / 'content' / 'manifest.json').read_text(encoding='utf-8'))
        self.assertEqual(manifest['config_input']['selection'], 'environment')
        self.assertEqual(manifest['shared_handover']['record']['gaps'], ['Partial crawl'])
        self.assertNotEqual(manifest['source'].get('crawl_id'), 'context-only')
        self.assertEqual(config.read_bytes(), original)
        self.assertEqual(usage.read_text(), '{"mcp_calls": 3}')

    def test_cli_rejects_other_site_exports(self):
        self.full_exports()
        result = self.cli(self.config(site='https://wrong.example/'))
        self.assertNotEqual(result.returncode, 0)
        self.assertIn('host', result.stderr.lower())
        self.assertFalse((self.path / 'run' / 'content' / 'findings.json').exists())

    def test_cli_rejects_mixed_site_filtered_exports(self):
        self.full_exports()
        self.write('h1_missing.csv', ['Address'], [['https://wrong.example/a']])
        self.assertNotEqual(self.cli(self.config()).returncode, 0)

    def test_cli_rejects_wrong_image_source_not_external_destination(self):
        self.full_exports()
        self.write('all_image_inlinks.csv', ['Source', 'Destination', 'Size (Bytes)'],
                   [['https://wrong.example/a', 'https://cdn.example.com/a.png', 204800]])
        self.assertNotEqual(self.cli(self.config()).returncode, 0)

    def test_export_override_clears_unverified_old_crawl_identity(self):
        self.full_exports()
        config = self.config(extra={'source': {'mode': 'saved_crawl', 'export_directory': 'old', 'crawl_file': 'old.seospider',
                                              'crawl_id': 'old', 'crawl_timestamp': '2000-01-01'}})
        script = ROOT / 'scripts' / 'onsite_checklist.py'
        result = subprocess.run([sys.executable, str(script), '--config', str(config), '--exports-dir', str(self.path),
                                 '--run-dir', str(self.path / 'override')], capture_output=True, text=True, encoding='utf-8')
        self.assertEqual(result.returncode, 0, result.stderr)
        manifest = json.loads((self.path / 'override' / 'content' / 'manifest.json').read_text(encoding='utf-8'))
        self.assertIsNone(manifest['crawl_timestamp'])
        self.assertEqual(manifest['source']['export_directory'], str(self.path.resolve()))
        self.assertIsNone(manifest['source'].get('crawl_id'))

    def test_saved_crawl_without_exports_requests_open_not_recrawl(self):
        config = self.config(extra={'source': {'mode': 'saved_crawl', 'crawl_file': 'example.seospider', 'allow_new_crawl': False}})
        result = self.cli(config)
        self.assertNotEqual(result.returncode, 0)
        self.assertIn('saved crawl', result.stderr.lower())

    def test_cli_disable_flow_does_not_read_exports(self):
        config = self.config(extra={'checks': {'content': False}})
        result = self.cli(config)
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertIn('disabled', result.stdout)
        self.assertFalse((self.path / 'run').exists())

    def test_cli_target_file_and_explicit_date(self):
        self.full_exports()
        self.write('page_titles_missing.csv', ['Address'], [['https://example.com/a']])
        self.write('targets.csv', ['Address'], [['https://example.com/a']])
        result = self.cli(self.config(extra={'content': {'target_urls_file': 'targets.csv'}}))
        self.assertEqual(result.returncode, 0, result.stderr)
        data = json.loads((self.path / 'run' / 'content' / 'findings.json').read_text(encoding='utf-8'))
        self.assertEqual(self.check(data, '17.1')['result'], '√')

    def test_cli_timezone_required_when_date_not_supplied(self):
        self.full_exports()
        result = self.cli(self.config(extra={'output': {'language': 'en'}}))
        self.assertNotEqual(result.returncode, 0)
        self.assertIn('timezone', result.stderr.lower())


if __name__ == '__main__':
    unittest.main()
