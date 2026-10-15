"""Render the onsite series content report, preserve literal text, verify read-back."""
from io import BytesIO
from pathlib import Path
import datetime
import re
import math
import unicodedata

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from extract import ITEMS

CHECK_HEADERS = ['Item No.', 'Item Name', 'Initial Check', 'Findings', 'Instruction']
GROUPS = (
    ('17', 'page_titles', None),
    ('18', 'meta_descriptions', None),
    ('19', 'headings', None),
    ('20', 'images', None),
)
DETAILS = {
    '17.1': ('17.1 Page title -Missing',
             [('Address', 'Address'), ('New Title', 'New Title'), ('Remarks', 'Remarks')]),
    '17.2': ('17.2 Page title - Duplicate',
             [('Address', 'Address'), ('Titles', 'Title 1'), ('New Title', 'New Title'),
              ('Instructions (If have)', 'Remarks')]),
    '17.3': ('17.3 Page title -Too Long',
             [('Address', 'Address'), ('Title - 1', 'Title 1'),
              ('Pixel Width ( max 580 pixels)', 'Title 1 Pixel Width'), ('Revised Title', 'Revised Title')]),
    '17.4': ('17.4 Page title - Multiple',
             [('Address', 'Address'), ('Title - 1', 'Title 1'), ('Title - 2', 'Title 2'), ('Instructions', 'Remarks')]),
    '18.1': ('18. Descriptions - Missing', [('Address', 'Address'), ('Remarks', 'Instruction')]),
    '18.2': ('18.2 Descriptions - Duplicate', [('Address', 'Address'), ('Meta Description', 'Meta Description 1')]),
    '18.3': ('18.3 Descriptions - Too Long',
             [('Address', 'Address'), ('Meta Description', 'Meta Description 1'), ('Pixel Width', 'Pixel Width')]),
    '19.1': ('19.1 H1 - Missing', [('Address', 'Address'), ('Remarks', 'Instruction')]),
    '19.2': ('19.2 H1 - Duplicate', [('Address', 'Address'), ('H1', 'H1-1'), ('Instructions', 'Remarks')]),
    '19.3': ('19.3 H1 - Multiple', [('Address', 'Address'), ('H1-1', 'H1-1'), ('H1-2', 'H1-2')]),
    '20.1': ('20.1 Image - Alt Text Missing',
             [('Source', 'Source'), ('Destination', 'Destination'), ('Example Screenshots (If Applicable)', 'Screenshots (If have)')]),
    '20.2': ('20.2 Image - Oversize', [('Source', 'Source'), ('Destination', 'Destination'), ('Size (Bytes)', 'Size (Bytes)')]),
}


def validate(data):
    datetime.date.fromisoformat(data['audit_date'])
    if data['audit_date'] != datetime.date.fromisoformat(data['audit_date']).isoformat():
        raise ValueError('audit_date must be YYYY-MM-DD')
    checks = data['checks']
    if len(checks) != 12 or {c['id'] for c in checks} != set(ITEMS):
        raise ValueError('Exactly the twelve content checks are required')
    for c in checks:
        name, sheet, headers = ITEMS[c['id']]
        if (c['item_name'], c['sheet'], c['headers']) != (name, sheet, headers):
            raise ValueError(f'Unexpected item schema: {c["id"]}')
        if c['result'] not in ('√', 'X', 'Human Check') or c['row_count'] != len(c['rows']):
            raise ValueError('Invalid result or row count')
        if c['rows'] and c['result'] != 'X':
            raise ValueError('Confirmed issue rows must be X')
        if not c['rows'] and c['result'] == 'X':
            raise ValueError('X requires confirmed issue rows')
        if c['gaps'] and c['result'] == '√':
            raise ValueError('Evidence gap cannot pass')
        if c['result'] == 'Human Check' and not c['gaps']:
            raise ValueError('Human Check requires an explicit evidence gap')
        if c['id'] == '18.3' and 'pixel_widths' in c:
            pixels = c['pixel_widths']
            if not isinstance(pixels, list) or len(pixels) != len(c['rows']) or any(
                    pixel is not None and (type(pixel) not in (int, float) or not math.isfinite(pixel) or pixel <= 0)
                    for pixel in pixels):
                raise ValueError('18.3 pixel_widths must align with rows and contain positive finite widths or null')
        for row in c['rows']:
            if len(row) > len(headers):
                raise ValueError('Too many detail columns')


def text(value):
    result = str(value if value is not None else '')
    if len(result.encode('utf-16-le', errors='surrogatepass')) // 2 > 32767:
        raise ValueError('Excel cell exceeds its length limit; keep full evidence in the archive')
    if re.search(r'[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]', result):
        raise ValueError('Unsupported Excel XML character')
    return result


def build_report(data, output_dir):
    validate(data)
    name = re.sub(r'[<>:"/\\|?*\x00-\x1f]', '_', data['site_name']).strip(' .')
    if not name or len(name.encode('utf-8')) > 180:
        raise ValueError('Provide a shorter portable site_name')
    paths = []
    for prefix, slug, _ in GROUPS:
        ids = sorted((item for item in ITEMS if item.split('.')[0] == prefix),
                     key=lambda item: tuple(map(int, item.split('.'))))
        item_range = f'{ids[0]}-{ids[-1]}'
        paths.append(Path(output_dir) / f'{item_range} {name}_{slug}_audit_{data["audit_date"]}.xlsx')
    for path in paths:
        if path.exists():
            raise FileExistsError(f'Report already exists: {path}; choose a new run directory')
    staged = []
    for prefix, _, sheet in GROUPS:
        checks = sorted((c for c in data['checks'] if c['id'].split('.')[0] == prefix),
                        key=lambda c: tuple(map(int, c['id'].split('.'))))
        expected = {'Checklist': [CHECK_HEADERS] +
                    [[c['id'], c['item_name'], c['result'], c['finding'], c['action']] for c in checks]}
        for c in checks:
            if c['result'] == 'X':
                detail_sheet, columns = DETAILS[c['id']]
                rows = [[label for label, _ in columns]]
                for index, row in enumerate(c['rows']):
                    values = dict(zip(c['headers'], row))
                    if c['id'] == '18.3':
                        values['Pixel Width'] = c.get('pixel_widths', [None] * len(c['rows']))[index]
                    rows.append([values.get(source, '') for _, source in columns])
                expected[detail_sheet] = rows
        staged.append((render_workbook(expected), expected))
    created = []
    try:
        Path(output_dir).mkdir(parents=True, exist_ok=True)
        for path, (contents, expected) in zip(paths, staged):
            with path.open('xb') as f:
                created.append(path)
                f.write(contents)
            verify_workbook(path, expected)
        return paths
    except Exception:
        for path in created:
            path.unlink(missing_ok=True)
        raise


def verify_workbook(source, expected):
    reopened = openpyxl.load_workbook(source)
    try:
        if reopened.sheetnames != list(expected):
            raise ValueError('Workbook worksheet verification failed')
        for sheet, rows in expected.items():
            actual = [[str(c.value) if c.value is not None else '' for c in row] for row in reopened[sheet].iter_rows()]
            wanted = [[text(v) for v in row] for row in rows]
            if actual != wanted or any(c.data_type == 'f' for row in reopened[sheet] for c in row):
                raise ValueError(f'Workbook content verification failed: {sheet}')
    finally:
        reopened.close()


def render_workbook(expected):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    try:
        for sheet, rows in expected.items():
            ws = wb.create_sheet(sheet)
            ws.freeze_panes = 'A2'
            ws.sheet_view.showGridLines = False
            widths = [12, 52, 20, 65, 60] if sheet == 'Checklist' else [
                12 if h == 'Item No.' else 22 if h in ('Title 1 Pixel Width', 'Pixel Width', 'Pixel Width ( max 580 pixels)',
                                                      'Character Count', 'Size (Bytes)') else 48 for h in rows[0]]
            for i, width in enumerate(widths, 1):
                ws.column_dimensions[get_column_letter(i)].width = width
            for r, values in enumerate(rows, 1):
                row_height = 30
                for col, value in enumerate(values, 1):
                    cell = ws.cell(r, col, text(value))
                    cell.data_type = 's'
                    cell.alignment = Alignment(wrap_text=True, vertical='top')
                    cell.font = Font(name='Arial', size=11, color='FFFFFF' if r == 1 else '203047', bold=r == 1)
                    fill = '183653' if r == 1 else ('F1F5FA' if r % 2 == 0 else 'FFFFFF')
                    if sheet == 'Checklist' and col == 3 and r > 1:
                        fill = {'√': 'DFF0E3', 'X': 'FADDDD', 'Human Check': 'FFF0C2'}[value]
                    cell.fill = PatternFill('solid', fgColor=fill)
                    units = sum(2 if unicodedata.east_asian_width(ch) in 'WF' else 1 for ch in text(value))
                    row_height = max(row_height, 15 * (1 + units // max(1, widths[col - 1] - 4)))
                ws.row_dimensions[r].height = min(409, row_height + 12)
            ws.auto_filter.ref = ws.dimensions
            ws.print_title_rows = '1:1'
            ws.sheet_properties.pageSetUpPr.fitToPage = True
            ws.page_setup.orientation = 'landscape'
            ws.page_setup.paperSize = ws.PAPERSIZE_A3
            ws.page_setup.fitToWidth = 1
            ws.page_setup.fitToHeight = 0

        buffer = BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        verify_workbook(buffer, expected)
        return buffer.getvalue()
    finally:
        wb.close()
